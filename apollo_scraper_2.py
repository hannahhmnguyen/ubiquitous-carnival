"""
Apollo.io scraper — finds exec contact info for companies in cleaned spreadsheet.
Usage:
    python3 apollo_scraper.py --input Womens_HRT_Hormones_Only_New_cleaned.xlsx --api-key YOUR_KEY
"""

import argparse
import time
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

APOLLO_BASE = "https://api.apollo.io/api/v1"

EXEC_TITLES = [
    "CEO", "Chief Executive Officer",
    "Founder", "Co-Founder",
    "President", "Owner",
    "Managing Director", "Medical Director",
    "Executive Director", "Partner",
]

def get_companies(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Deduped"]
    rows = list(ws.iter_rows(values_only=True))
    companies = []
    for r in rows[1:]:
        companies.append({
            "company":           r[0] or "",
            "existing_first":    str(r[1]) if r[1] else "",
            "existing_last":     str(r[2]) if r[2] else "",
            "website":           r[3] or "",
            "existing_phone":    str(r[4]) if r[4] else "",
            "existing_title":    str(r[5]) if r[5] else "",
            "existing_email":    str(r[6]) if r[6] else "",
            "existing_linkedin": str(r[7]) if r[7] else "",
            "reviews":           r[8],
        })
    return companies


def clean_domain(website):
    if not website:
        return None
    w = str(website).strip()
    w = w.replace("https://", "").replace("http://", "").replace("www.", "")
    return w.split("/")[0] or None


def search_person(api_key, company_name, domain):
    headers = {"X-Api-Key": api_key, "Content-Type": "application/json"}
    payload = {
        "person_titles": EXEC_TITLES,
        "page": 1,
        "per_page": 5,
    }
    if domain:
        payload["q_organization_domains"] = domain
    else:
        payload["q_keywords"] = company_name

    resp = requests.post(f"{APOLLO_BASE}/mixed_people/api_search", json=payload, headers=headers, timeout=30)
    if resp.status_code == 429:
        print("  Rate limited, waiting 60s...")
        time.sleep(60)
        resp = requests.post(f"{APOLLO_BASE}/mixed_people/api_search", json=payload, headers=headers, timeout=30)
    if resp.status_code != 200:
        print(f"  API error {resp.status_code}: {resp.text[:200]}")
        return None

    people = resp.json().get("people", [])
    if not people:
        return None

    priority = ["ceo", "founder", "president", "owner", "director"]
    for person in people:
        if any(kw in (person.get("title") or "").lower() for kw in priority):
            return person
    return people[0]


def extract_contact(person):
    email = person.get("email") or ""
    phones = person.get("phone_numbers", [])
    phone = phones[0].get("sanitized_number", "") if phones else ""
    linkedin = person.get("linkedin_url") or ""
    return email, phone, linkedin


def run(input_path, api_key, output_path, batch_size):
    companies = get_companies(input_path)
    total = len(companies)
    print(f"Found {total} companies in spreadsheet")
    print(f"Running batch of {batch_size}\n")

    batch = companies[:batch_size]
    results = []

    for i, co in enumerate(batch):
        name = co["company"]
        domain = clean_domain(co["website"])
        print(f"[{i+1}/{len(batch)}] {name} ({domain})")

        if co["existing_first"] and co["existing_email"]:
            print("  Already complete — skipping")
            results.append({
                **co,
                "first": co["existing_first"],
                "last": co["existing_last"],
                "email": co["existing_email"],
                "phone": co["existing_phone"],
                "title": co["existing_title"],
                "linkedin": co["existing_linkedin"],
            })
            continue

        person = search_person(api_key, name, domain)
        if not person:
            print("  No results found")
            results.append({
                **co,
                "first": co["existing_first"],
                "last": co["existing_last"],
                "email": co["existing_email"],
                "phone": co["existing_phone"],
                "title": co["existing_title"],
                "linkedin": co["existing_linkedin"],
            })
            time.sleep(1)
            continue

        first = person.get("first_name") or ""
        last = person.get("last_name") or ""
        title = person.get("title") or ""
        email, phone, linkedin = extract_contact(person)

        print(f"  Found: {first} {last} | {title} | {email} | {phone}")

        results.append({
            **co,
            "first":   first   or co["existing_first"],
            "last":    last    or co["existing_last"],
            "email":   email   or co["existing_email"],
            "phone":   phone   or co["existing_phone"],
            "title":   title   or co["existing_title"],
            "linkedin": linkedin or co["existing_linkedin"],
        })
        time.sleep(1.2)

    # Write output
    wb_out = openpyxl.Workbook()
    ws = wb_out.active
    ws.title = "Exec Info"

    headers = ["Company Name", "Exec First Name", "Exec Last Name", "Website",
               "Phone Number", "Exec Title", "Exec Email", "Exec LinkedIn", "Reviews"]

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r in results:
        ws.append([
            r["company"], r.get("first",""), r.get("last",""),
            r["website"], r.get("phone",""), r.get("title",""),
            r.get("email",""), r.get("linkedin",""), r["reviews"],
        ])

    col_widths = [35, 15, 15, 35, 18, 30, 35, 50, 10]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    wb_out.save(output_path)
    print(f"\nDone! Saved {len(results)} rows to {output_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--api-key", default="DzmJf2rVwVn4rysfBMF6Ow")
    parser.add_argument("--output", default="Exec_Contacts.xlsx")
    parser.add_argument("--batch", type=int, default=5)
    args = parser.parse_args()
    run(args.input, args.api_key, args.output, args.batch)
