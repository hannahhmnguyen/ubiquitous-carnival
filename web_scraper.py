"""
Web scraper — visits clinic websites and extracts contact info directly.
Usage:
    python3 web_scraper.py --input Womens_HRT_Hormones_Only_New_cleaned.xlsx
    python3 web_scraper.py --input Womens_HRT_Hormones_Only_New_cleaned.xlsx --batch 20 --output results.xlsx
"""

import argparse
import re
import time
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from bs4 import BeautifulSoup
from urllib.parse import urljoin, urlparse

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/120.0.0.0 Safari/537.36"
    )
}

CONTACT_PAGE_HINTS = [
    "contact", "about", "team", "staff", "our-team", "about-us",
    "meet-the-team", "providers", "meet-us", "who-we-are",
]

EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}")
PHONE_RE = re.compile(r"(\(?\d{3}\)?[\s.\-]?\d{3}[\s.\-]?\d{4})")

NAME_TITLE_KEYWORDS = [
    "md", "do", "np", "pa", "rn", "fnp", "aprn",
    "founder", "owner", "director", "ceo", "president", "physician",
    "doctor", "dr.", "dr ",
]


def get_companies(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Deduped"]
    rows = list(ws.iter_rows(values_only=True))
    companies = []
    for r in rows[1:]:
        companies.append({
            "company":           r[0] or "",
            "existing_first":    str(r[1]) if r[1] else "",
            "existing_last":     str(r[2]) if r[2] else "",
            "website":           str(r[3]).strip() if r[3] else "",
            "existing_phone":    str(r[4]) if r[4] else "",
            "existing_title":    str(r[5]) if r[5] else "",
            "existing_email":    str(r[6]) if r[6] else "",
            "existing_linkedin": str(r[7]) if r[7] else "",
            "reviews":           r[8],
        })
    return companies


def normalize_url(website):
    if not website:
        return None
    w = website.strip()
    if not w.startswith("http"):
        w = "https://" + w
    return w.rstrip("/")


def fetch(url, timeout=15):
    try:
        resp = requests.get(url, headers=HEADERS, timeout=timeout, allow_redirects=True)
        if resp.status_code == 200:
            return resp.text
    except Exception:
        pass
    # retry with http if https failed
    if url.startswith("https://"):
        try:
            resp = requests.get(url.replace("https://", "http://"), headers=HEADERS, timeout=timeout, allow_redirects=True)
            if resp.status_code == 200:
                return resp.text
        except Exception:
            pass
    return None


def find_contact_pages(base_url, soup):
    """Return a list of candidate contact/about page URLs found in nav links."""
    candidates = []
    for a in soup.find_all("a", href=True):
        href = a["href"].lower()
        text = a.get_text(strip=True).lower()
        if any(h in href or h in text for h in CONTACT_PAGE_HINTS):
            full = urljoin(base_url, a["href"])
            # keep same domain only
            if urlparse(full).netloc == urlparse(base_url).netloc:
                if full not in candidates:
                    candidates.append(full)
    return candidates[:6]


def extract_emails(text):
    found = EMAIL_RE.findall(text)
    # filter out image/asset false positives
    return [e for e in found if not any(e.endswith(x) for x in [".png", ".jpg", ".gif", ".svg"])]


def extract_phones(text):
    found = PHONE_RE.findall(text)
    seen = set()
    out = []
    for p in found:
        clean = re.sub(r"\D", "", p)
        if clean not in seen and len(clean) == 10:
            seen.add(clean)
            out.append(p.strip())
    return out


def extract_name_from_soup(soup):
    """Best-effort: find a person name near a title keyword."""
    text = soup.get_text(separator="\n")
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    for i, line in enumerate(lines):
        ll = line.lower()
        if any(kw in ll for kw in NAME_TITLE_KEYWORDS):
            # look at surrounding lines for a name-like string (2-4 words, title case)
            for candidate in lines[max(0, i-2):i+3]:
                words = candidate.split()
                if 2 <= len(words) <= 4 and all(w[0].isupper() for w in words if w[0].isalpha()):
                    # skip lines that are just titles/keywords
                    if not any(kw in candidate.lower() for kw in ["clinic", "center", "medical", "wellness", "health"]):
                        return candidate
    return ""


def scrape_site(base_url):
    """Scrape homepage + contact/about pages. Returns (emails, phones, name)."""
    html = fetch(base_url)
    if not html:
        return [], [], ""

    soup = BeautifulSoup(html, "html.parser")
    emails = extract_emails(html)
    phones = extract_phones(html)
    name = extract_name_from_soup(soup)

    contact_pages = find_contact_pages(base_url, soup)
    for url in contact_pages:
        if emails and phones and name:
            break
        page_html = fetch(url)
        if not page_html:
            continue
        page_soup = BeautifulSoup(page_html, "html.parser")
        emails += extract_emails(page_html)
        phones += extract_phones(page_html)
        if not name:
            name = extract_name_from_soup(page_soup)
        time.sleep(0.5)

    # deduplicate
    emails = list(dict.fromkeys(emails))
    phones = list(dict.fromkeys(phones))
    return emails, phones, name


def run(input_path, output_path, batch_size):
    companies = get_companies(input_path)
    total = len(companies)
    print(f"Found {total} companies in spreadsheet")
    print(f"Running batch of {batch_size}\n")

    batch = companies[:batch_size]
    results = []

    for i, co in enumerate(batch):
        name = co["company"]
        url = normalize_url(co["website"])
        print(f"[{i+1}/{len(batch)}] {name} ({url})")

        if co["existing_first"] and co["existing_email"]:
            print("  Already complete — skipping")
            results.append({**co,
                "first": co["existing_first"], "last": co["existing_last"],
                "email": co["existing_email"], "phone": co["existing_phone"],
                "title": co["existing_title"], "linkedin": co["existing_linkedin"],
            })
            continue

        if not url:
            print("  No website — skipping")
            results.append({**co,
                "first": co["existing_first"], "last": co["existing_last"],
                "email": co["existing_email"], "phone": co["existing_phone"],
                "title": co["existing_title"], "linkedin": co["existing_linkedin"],
            })
            continue

        emails, phones, found_name = scrape_site(url)

        email = emails[0] if emails else co["existing_email"]
        phone = phones[0] if phones else co["existing_phone"]

        # split found_name into first/last
        first, last = co["existing_first"], co["existing_last"]
        if found_name and not first:
            parts = found_name.split()
            first = parts[0]
            last = " ".join(parts[1:]) if len(parts) > 1 else ""

        print(f"  Email: {email or '—'}  Phone: {phone or '—'}  Name: {found_name or '—'}")

        results.append({**co,
            "first": first, "last": last,
            "email": email, "phone": phone,
            "title": co["existing_title"],
            "linkedin": co["existing_linkedin"],
        })
        time.sleep(1)

    # Write output
    wb_out = openpyxl.Workbook()
    ws = wb_out.active
    ws.title = "Exec Info"

    col_headers = ["Company Name", "Exec First Name", "Exec Last Name", "Website",
                   "Phone Number", "Exec Title", "Exec Email", "Exec LinkedIn", "Reviews"]

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    ws.append(col_headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r in results:
        ws.append([
            r["company"], r.get("first", ""), r.get("last", ""),
            r["website"], r.get("phone", ""), r.get("title", ""),
            r.get("email", ""), r.get("linkedin", ""), r["reviews"],
        ])

    col_widths = [35, 15, 15, 35, 18, 30, 35, 50, 10]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    wb_out.save(output_path)
    print(f"\nDone! Saved {len(results)} rows to {output_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", default="Exec_Contacts.xlsx")
    parser.add_argument("--batch", type=int, default=5)
    args = parser.parse_args()
    run(args.input, args.output, args.batch)
